import os
import sys
import mammoth

def convert_to_markdown(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    output_path = file_path.rsplit('.', 1)[0] + '.md'
    
    try:
        if ext == '.docx':
            with open(file_path, "rb") as docx_file:
                result = mammoth.convert_to_markdown(docx_file)
                markdown = result.value
                with open(output_path, "w") as md_file:
                    md_file.write(markdown)
            print(f"Converted {file_path} to {output_path}")
            return output_path
        elif ext == '.doc':
            # .doc requires different handling
            print(f"WARNING: .doc format found for {file_path}. Please save as .docx for automated processing.")
            return None
        elif ext == '.pptx':
            try:
                from pptx import Presentation
                from pptx.enum.shapes import MSO_SHAPE_TYPE
            except ImportError:
                print("ERROR: python-pptx package is not installed. Please run: pip3 install python-pptx")
                return None
            
            prs = Presentation(file_path)
            markdown_content = []
            markdown_content.append(f"# Presentation: {os.path.basename(file_path)}\n\n")
            
            # Determine screenshots directory
            project_root = os.path.abspath(os.path.join(os.path.dirname(file_path), ".."))
            screenshot_dir = os.path.join(project_root, "ui_analysis", "screenshots")
            if not os.path.isdir(screenshot_dir):
                # Fallback to a folder inside discovery
                screenshot_dir = os.path.join(os.path.dirname(file_path), "extracted_images")
            os.makedirs(screenshot_dir, exist_ok=True)
            
            presentation_name = os.path.splitext(os.path.basename(file_path))[0]
            
            for idx, slide in enumerate(prs.slides):
                markdown_content.append(f"## Slide {idx + 1}\n\n")
                
                # Check for speaker notes
                notes_text = ""
                if slide.notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                
                # Extract shapes
                image_count = 0
                for shape in slide.shapes:
                    # Check if it has a text frame
                    if shape.has_text_frame:
                        text = shape.text_frame.text.strip()
                        if text:
                            # Format title shape as h3
                            is_title = False
                            try:
                                if slide.shapes.title == shape:
                                    is_title = True
                            except Exception:
                                pass
                            
                            if is_title:
                                markdown_content.append(f"### {text}\n\n")
                            else:
                                markdown_content.append(f"{text}\n\n")
                                
                    # Check if it is a table
                    elif shape.has_table:
                        table = shape.table
                        rows_data = []
                        for row in table.rows:
                            row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                            rows_data.append(row_cells)
                        
                        if rows_data:
                            headers = rows_data[0]
                            markdown_content.append("| " + " | ".join(headers) + " |")
                            markdown_content.append("| " + " | ".join(["---"] * len(headers)) + " |")
                            for row in rows_data[1:]:
                                markdown_content.append("| " + " | ".join(row) + " |")
                            markdown_content.append("\n\n")
                            
                    # Check if it is a picture (embedded screenshot/mockup)
                    else:
                        is_picture = False
                        try:
                            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                                is_picture = True
                        except Exception:
                            pass
                        
                        if not is_picture:
                            try:
                                if hasattr(shape, "image") and shape.image is not None:
                                    is_picture = True
                            except Exception:
                                pass
                        
                        if is_picture:
                            try:
                                image = shape.image
                                image_bytes = image.blob
                                ext_name = image.ext if image.ext else "png"
                                
                                # Format filename
                                image_filename = f"{presentation_name}_slide_{idx + 1}_image_{image_count}.{ext_name}"
                                image_file_path = os.path.join(screenshot_dir, image_filename)
                                
                                # Write file
                                with open(image_file_path, "wb") as f:
                                    f.write(image_bytes)
                                
                                # Generate absolute path link
                                abs_path = os.path.abspath(image_file_path)
                                if abs_path.startswith('/'):
                                    file_url = f"file://{abs_path}"
                                else:
                                    file_url = f"file:///{abs_path.replace('\\', '/')}"
                                
                                markdown_content.append(f"![Slide {idx + 1} Mockup Image {image_count}]({file_url})\n\n")
                                image_count += 1
                            except Exception as img_err:
                                print(f"Error extracting image on slide {idx + 1}: {img_err}")
                
                # Append speaker notes at the end of the slide if present
                if notes_text:
                    markdown_content.append("> **Speaker Notes:**\n")
                    for line in notes_text.split('\n'):
                        markdown_content.append(f"> {line}\n")
                    markdown_content.append("\n")
                
                markdown_content.append("---\n\n")
                
            with open(output_path, "w", encoding="utf-8") as md_file:
                md_file.write("".join(markdown_content))
            print(f"Converted {file_path} to {output_path}")
            return output_path
        elif ext == '.ppt':
            print(f"WARNING: .ppt format found for {file_path}. Please save as .pptx for automated processing.")
            return None

    except Exception as e:
        print(f"Error converting {file_path}: {e}")
        return None

if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "./discovery"
    if os.path.isfile(target):
        convert_to_markdown(target)
    elif os.path.isdir(target):
        for f in os.listdir(target):
            if f.endswith(('.docx', '.doc', '.pptx', '.ppt')):
                convert_to_markdown(os.path.join(target, f))
