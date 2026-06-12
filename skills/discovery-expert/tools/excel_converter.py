import os
import sys
import json
import re

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, str):
        return val.strip().replace('\xa0', ' ').replace('\u00a0', ' ')
    return val

def clean_path(path):
    if not path:
        return "/unknown"
    path = path.replace("{endPoint}", "").replace("{endpoint}", "")
    path = path.strip()
    if not path.startswith("/"):
        path = "/" + path
    return path

def clean_and_parse_json(json_str):
    if not json_str:
        return None
    cleaned = json_str.strip()
    cleaned = cleaned.replace('“', '"').replace('”', '"').replace('’', '"').replace('‘', '"')
    cleaned = cleaned.replace('\xa0', ' ').replace('\u00a0', ' ')
    # Remove trailing commas inside JSON objects/arrays
    cleaned = re.sub(r',\s*([\]}])', r'\1', cleaned)
    try:
        return json.loads(cleaned)
    except Exception:
        try:
            import ast
            return ast.literal_eval(cleaned)
        except Exception:
            return cleaned

def json_to_schema(data):
    if isinstance(data, dict):
        properties = {}
        for k, v in data.items():
            properties[k] = json_to_schema(v)
        return {
            "type": "object",
            "properties": properties
        }
    elif isinstance(data, list):
        if len(data) > 0:
            items_schema = json_to_schema(data[0])
        else:
            items_schema = {"type": "string"}
        return {
            "type": "array",
            "items": items_schema
        }
    elif isinstance(data, bool):
        return {"type": "boolean"}
    elif isinstance(data, int):
        return {"type": "integer"}
    elif isinstance(data, float):
        return {"type": "number"}
    elif data is None:
        return {"type": "string", "nullable": True}
    else:
        return {"type": "string"}

def convert_xlsx_to_openapi(file_path, output_dir=None):
    try:
        import openpyxl
        import yaml
    except ImportError:
        print("ERROR: openpyxl and pyyaml are required. Please run: pip3 install openpyxl pyyaml")
        return None

    print(f"Loading workbook {file_path}...")
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        print(f"Error opening workbook {file_path}: {e}")
        return None

    filename = os.path.splitext(os.path.basename(file_path))[0]
    if output_dir is None:
        # Save to discovery/specs/
        project_root = os.path.abspath(os.path.join(os.path.dirname(file_path), ".."))
        output_dir = os.path.join(project_root, "discovery", "specs")
        if not os.path.isdir(output_dir):
            output_dir = os.path.dirname(file_path)
            
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f"{filename}.yaml")

    openapi_spec = {
        "openapi": "3.0.0",
        "info": {
            "title": f"API Specification - {filename.replace('_', ' ')}",
            "description": f"Generated from {os.path.basename(file_path)}",
            "version": "1.0.0"
        },
        "servers": [
            {"url": "https://api.example.com", "description": "Mock Server"}
        ],
        "paths": {}
    }

    # Search for server urls in index sheet or metadata if available
    # For now we use standard default server or extract from sheets

    has_valid_sheets = False
    for sheet_name in wb.sheetnames:
        # Skip index sheets or other non-endpoint sheets
        if sheet_name in ['Index', 'Mapping_Role', 'Sheet1', 'Sheet2']:
            continue
            
        sheet = wb[sheet_name]
        
        # Determine if this sheet actually specifies an API by checking cells
        # We look for "API Name" or "API Code"
        is_api_sheet = False
        for r in list(sheet.iter_rows(values_only=True))[:10]:
            r_str = " ".join([str(c) for c in r if c is not None])
            if "API Name" in r_str or "API Code" in r_str or "Endpoint" in r_str:
                is_api_sheet = True
                break
                
        if not is_api_sheet:
            continue

        has_valid_sheets = True
        print(f"  Processing endpoint sheet: {sheet_name}...")
        
        api_name = ""
        api_code = ""
        method = "POST"
        endpoint = ""
        ui_component = ""
        
        request_fields = []
        request_examples = {}
        response_examples = {}
        
        state = "META"
        current_case = ""
        
        # Iterate over all rows in sheet
        for row in sheet.iter_rows(values_only=True):
            cells = [clean_val(c) for c in row]
            if not any(cells):
                continue
                
            first_cell = cells[0]
            
            # Check state transitions
            if first_cell == "Request Body" or (first_cell == "No" and "field Name" in cells):
                state = "REQ_FIELDS"
                continue
            elif first_cell == "Request Body  [json]" or first_cell == "Request Body [json]":
                state = "REQ_JSON"
                continue
            elif first_cell == "Response Body [json]" or first_cell == "Response Body  [json]":
                state = "RESP_JSON"
                continue
                
            if state == "META":
                for i, cell in enumerate(cells):
                    cell_str = str(cell)
                    if "API Name" in cell_str:
                        api_name = next((cells[j] for j in range(i+1, len(cells)) if cells[j]), "")
                    elif "API Code" in cell_str:
                        api_code = next((cells[j] for j in range(i+1, len(cells)) if cells[j]), "")
                    elif "Method" in cell_str:
                        method = next((cells[j] for j in range(i+1, len(cells)) if cells[j]), "").upper()
                    elif "Endpoint" in cell_str:
                        endpoint = next((cells[j] for j in range(i+1, len(cells)) if cells[j]), "")
                    elif "UI Component" in cell_str:
                        ui_component = next((cells[j] for j in range(i+1, len(cells)) if cells[j]), "")
                        
            elif state == "REQ_FIELDS":
                if "field Name" in cells or first_cell == "No":
                    continue
                field_name = cells[1]
                if field_name:
                    mandatory = cells[2]
                    type_str = cells[3]
                    example_val = cells[4]
                    request_fields.append({
                        "name": field_name,
                        "mandatory": "Y" in mandatory.upper(),
                        "type": type_str,
                        "example": example_val
                    })
                    
            elif state == "REQ_JSON":
                json_str = ""
                label = ""
                for i, cell in enumerate(cells):
                    if "{" in cell and "}" in cell:
                        json_str = cell
                    elif cell.startswith("##"):
                        label = cell.replace("##", "").strip()
                
                if json_str:
                    parsed_json = clean_and_parse_json(json_str)
                    if parsed_json:
                        if not label:
                            label = f"example_{len(request_examples) + 1}"
                        request_examples[label] = parsed_json
                        
            elif state == "RESP_JSON":
                case_match = None
                for cell in cells:
                    if "Case" in cell:
                        case_match = cell.replace("Case", "").strip()
                        break
                
                if case_match:
                    current_case = case_match
                    continue
                    
                json_str = ""
                for cell in cells:
                    if "{" in cell and "}" in cell:
                        json_str = cell
                        break
                        
                if json_str:
                    parsed_json = clean_and_parse_json(json_str)
                    if parsed_json:
                        label = current_case if current_case else "Success"
                        response_examples[label] = parsed_json
                        current_case = ""

        # Construct path item
        path_str = clean_path(endpoint)
        method_lower = method.lower() if method else "post"
        
        if path_str not in openapi_spec["paths"]:
            openapi_spec["paths"][path_str] = {}
            
        # Build schema for request
        req_properties = {}
        req_required = []
        
        for f in request_fields:
            f_type = f["type"].lower()
            oas_type = "string"
            if "number" in f_type or "double" in f_type or "float" in f_type:
                oas_type = "number"
            elif "integer" in f_type or "int" in f_type:
                oas_type = "integer"
            elif "boolean" in f_type or "bool" in f_type:
                oas_type = "boolean"
            elif "array" in f_type:
                oas_type = "array"
                
            prop = {"type": oas_type}
            if f["example"]:
                prop["example"] = f["example"]
                
            req_properties[f["name"]] = prop
            if f["mandatory"]:
                req_required.append(f["name"])
                
        if not req_properties and request_examples:
            first_ex = list(request_examples.values())[0]
            req_schema = json_to_schema(first_ex)
        else:
            req_schema = {
                "type": "object",
                "properties": req_properties
            }
            if req_required:
                req_schema["required"] = req_required
                
        # Build requestBody examples
        req_body_obj = {}
        if req_schema or request_examples:
            req_body_obj = {
                "required": True,
                "content": {
                    "application/json": {
                        "schema": req_schema
                    }
                }
            }
            if request_examples:
                req_body_obj["content"]["application/json"]["examples"] = {
                    lbl: {"value": val} for lbl, val in request_examples.items()
                }
                
        # Build response schema and examples
        resp_schema = {"type": "object"}
        if response_examples:
            first_resp = list(response_examples.values())[0]
            resp_schema = json_to_schema(first_resp)
            
        responses_obj = {
            "200": {
                "description": "Successful Response",
                "content": {
                    "application/json": {
                        "schema": resp_schema
                    }
                }
            }
        }
        
        if response_examples:
            responses_obj["200"]["content"]["application/json"]["examples"] = {
                lbl: {"value": val} for lbl, val in response_examples.items()
            }
            
        # Assemble operation
        operation = {
            "summary": api_name if api_name else sheet_name,
            "operationId": api_code if api_code else sheet_name,
            "description": f"UI Component: {ui_component}" if ui_component else "",
            "responses": responses_obj
        }
        
        if req_body_obj:
            operation["requestBody"] = req_body_obj
            
        openapi_spec["paths"][path_str][method_lower] = operation

    if not has_valid_sheets:
        print(f"WARNING: No valid API endpoint sheets found in {file_path}")
        return None

    # Write to output file
    with open(output_path, "w", encoding="utf-8") as f:
        yaml.dump(openapi_spec, f, default_flow_style=False, allow_unicode=True, sort_keys=False)

    print(f"Successfully created single OpenAPI spec file: {output_path}")
    return output_path

if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "./discovery"
    if os.path.isfile(target):
        if target.endswith(('.xlsx', '.xls')):
            convert_xlsx_to_openapi(target)
    elif os.path.isdir(target):
        for f in os.listdir(target):
            if f.endswith(('.xlsx', '.xls')):
                convert_xlsx_to_openapi(os.path.join(target, f))
